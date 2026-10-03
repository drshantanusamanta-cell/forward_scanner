"""
Forward Sentiment Scanner — NSE watchlist
Forward-looking analyst data from yfinance (G1–G4 of the Forward Sentiment Matrix).

Run:  streamlit run app.py
Data: yfinance (unofficial Yahoo Finance API). Coverage is limited to stocks that
      analysts follow; small caps will often show "no forward data".
Not investment advice.
"""
from __future__ import annotations

import datetime as dt
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import numpy as np
import pandas as pd
import streamlit as st
import yfinance as yf

APP_DIR = Path(__file__).parent
STOCKS_FILE = APP_DIR / "stocks.csv"
SNAP_DIR = APP_DIR / "snapshots"
SNAP_DIR.mkdir(exist_ok=True)

SYMBOL_COLS = ("symbol", "ticker", "nse code", "nse_code", "nsecode", "stock", "tradingsymbol", "tradingview_symbol")

# Matrix defaults: (label, direction, bull, bear, weight, value<=0 scores -1)
DEFAULT_RULES = {
    "G1": ("Analyst target upside %", "Higher", 15.0, 0.0, 5, False),
    "G2": ("Forward PE ÷ TTM PE", "Lower", 0.85, 1.0, 3, True),
    "G3": ("PEG (fwd PE ÷ next-FY EPS growth)", "Lower", 1.0, 2.0, 4, True),
    "G4": ("Next-FY EPS revision, 90d %", "Higher", 3.0, -3.0, 4, False),
}


# ----------------------------------------------------------------------------- helpers
def _num(x) -> float:
    try:
        x = float(x)
        return x if np.isfinite(x) else np.nan
    except (TypeError, ValueError):
        return np.nan


def load_universe(src) -> pd.DataFrame:
    """Read the watchlist CSV/TXT. Returns columns: yf (Yahoo ticker), symbol, section.

    Accepts a 'symbol' column (plus optional 'exchange' and 'section'), a TradingView
    'NSE:XYZ' column, or a header-less one-symbol-per-line file.
    """
    df = pd.read_csv(src, header=0, dtype=str)
    cols = {c.strip().lower(): c for c in df.columns}
    col = next((cols[c] for c in SYMBOL_COLS if c in cols), None)
    if col is None:  # no recognised header: first column holds the symbols
        if hasattr(src, "seek"):
            src.seek(0)
        df = pd.read_csv(src, header=None, dtype=str)
        cols, col = {}, df.columns[0]
    out = pd.DataFrame({"raw": df[col].astype(str).str.strip().str.upper()})
    out["section"] = df[cols["section"]].fillna("").str.strip() if "section" in cols else ""
    exch = df[cols["exchange"]].fillna("NSE").str.strip().str.upper() if "exchange" in cols else "NSE"
    out["exch"] = exch
    out = out[(out["raw"] != "") & (out["raw"] != "NAN") & (~out["raw"].str.startswith("#"))]
    pref = out["raw"].str.extract(r"^(NSE|BSE):", expand=False)
    out["exch"] = pref.fillna(out["exch"])
    out["symbol"] = (out["raw"].str.replace(r"^(NSE|BSE):", "", regex=True)
                     .str.replace(r"\.(NS|BO)$", "", regex=True)
                     .str.replace("_", "-", regex=False))  # TradingView NAM_INDIA -> Yahoo NAM-INDIA
    out["yf"] = out["symbol"] + np.where(out["exch"] == "BSE", ".BO", ".NS")
    out = (out.groupby("yf", sort=False)
              .agg(symbol=("symbol", "first"),
                   section=("section", lambda x: "; ".join(dict.fromkeys(v for v in x if v))))
              .reset_index())
    return out[["yf", "symbol", "section"]]


def fetch_one(sym: str, retries: int = 2) -> dict:
    """Pull analyst/forward fields for one ticker. Never raises."""
    last_err = ""
    for attempt in range(retries + 1):
        try:
            tk = yf.Ticker(sym)
            info = tk.info or {}
            row = {
                "symbol": sym.rsplit(".", 1)[0],
                "name": info.get("shortName") or info.get("longName"),
                "sector": info.get("sector"),
                "price": _num(info.get("currentPrice") or info.get("regularMarketPrice")),
                "mcap_cr": _num(info.get("marketCap")) / 1e7,
                "analysts": _num(info.get("numberOfAnalystOpinions")),
                "target_mean": _num(info.get("targetMeanPrice")),
                "target_median": _num(info.get("targetMedianPrice")),
                "target_high": _num(info.get("targetHighPrice")),
                "target_low": _num(info.get("targetLowPrice")),
                "rec_mean": _num(info.get("recommendationMean")),
                "rec_key": info.get("recommendationKey"),
                "eps_ttm": _num(info.get("trailingEps")),
                "eps_fwd": _num(info.get("forwardEps")),
            }
            # EPS trend: consensus now vs 30/90 days ago
            try:
                et = tk.eps_trend
            except Exception:
                et = None
            if et is not None and not et.empty:
                if "+1y" in et.index:
                    r = et.loc["+1y"]
                    row["eps_ny_now"] = _num(r.get("current"))
                    row["eps_ny_30d"] = _num(r.get("30daysAgo"))
                    row["eps_ny_90d"] = _num(r.get("90daysAgo"))
                if "0y" in et.index:
                    row["eps_cy_now"] = _num(et.loc["0y"].get("current"))
            # Revision counts (yfinance counts look inconsistent; shown, not scored)
            try:
                er = tk.eps_revisions
            except Exception:
                er = None
            if er is not None and not er.empty and "+1y" in er.index:
                r = er.loc["+1y"]
                row["up_30d"] = _num(r.get("upLast30days"))
                row["down_30d"] = _num(r.get("downLast30days"))
            row["fetched"] = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
            row["error"] = "" if pd.notna(row["price"]) else "No price returned (check symbol)"
            return row
        except Exception as e:  # network / rate limit / parsing
            last_err = f"{type(e).__name__}: {str(e)[:120]}"
            time.sleep(2 * (attempt + 1))
    return {"symbol": sym.rsplit(".", 1)[0], "error": last_err}


def fetch_all(symbols: list[str], workers: int, progress=None) -> pd.DataFrame:
    rows, done = [], 0
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futs = {ex.submit(fetch_one, s): s for s in symbols}
        for f in as_completed(futs):
            rows.append(f.result())
            done += 1
            if progress:
                progress.progress(done / len(symbols), text=f"Fetched {done}/{len(symbols)}: {futs[f]}")
    df = pd.DataFrame(rows)
    order = {s.rsplit(".", 1)[0]: i for i, s in enumerate(symbols)}
    return df.sort_values("symbol", key=lambda c: c.map(order)).reset_index(drop=True)


def score_value(v, direction, bull, bear, neg_is_bad) -> float:
    if pd.isna(v):
        return np.nan
    if neg_is_bad and v <= 0:
        return -1.0
    if direction == "Higher":
        return 1.0 if v >= bull else (-1.0 if v <= bear else 0.0)
    return 1.0 if v <= bull else (-1.0 if v >= bear else 0.0)


def score_relative(v: pd.Series, direction: str, neg_is_bad: bool) -> pd.Series:
    """Rank within the watchlist: best third +1, middle 0, worst third -1.
    Removes the market-wide optimism in analyst numbers (most targets sit above price)."""
    out = pd.Series(np.nan, index=v.index)
    bad = (v <= 0) if neg_is_bad else pd.Series(False, index=v.index)
    ok = v.notna() & ~bad
    if ok.sum() >= 3:
        pct = v[ok].rank(pct=True, ascending=(direction == "Higher"))
        out[ok] = np.where(pct > 2 / 3, 1.0, np.where(pct <= 1 / 3, -1.0, 0.0))
    out[bad & v.notna()] = -1.0
    return out


def compute(raw: pd.DataFrame, rules: dict, min_analysts: int, mode: str = "Relative") -> pd.DataFrame:
    df = raw.copy()
    for c in ["price", "target_mean", "eps_ttm", "eps_fwd", "eps_ny_now", "eps_ny_30d",
              "eps_ny_90d", "eps_cy_now", "analysts", "up_30d", "down_30d"]:
        if c not in df:
            df[c] = np.nan
    covered = df["analysts"].fillna(0) >= min_analysts
    df["covered"] = covered

    pe_ttm = np.where(df["eps_ttm"] > 0, df["price"] / df["eps_ttm"], np.nan)
    pe_fwd = df["price"] / df["eps_fwd"]
    df["pe_ttm"] = pe_ttm
    df["pe_fwd"] = pe_fwd
    growth = np.where(df["eps_cy_now"] > 0, (df["eps_ny_now"] / df["eps_cy_now"] - 1) * 100, np.nan)
    df["eps_growth_ny"] = growth

    df["G1"] = (df["target_mean"] / df["price"] - 1) * 100
    df["G2"] = np.where(np.isfinite(pe_ttm), pe_fwd / pe_ttm, np.nan)
    df["G3"] = np.where(np.isfinite(growth) & (growth != 0), pe_fwd / growth, np.nan)
    df["G4"] = np.where(df["eps_ny_90d"] > 0, (df["eps_ny_now"] / df["eps_ny_90d"] - 1) * 100, np.nan)
    df["G4_30d"] = np.where(df["eps_ny_30d"] > 0, (df["eps_ny_now"] / df["eps_ny_30d"] - 1) * 100, np.nan)
    tot = df["up_30d"].fillna(0) + df["down_30d"].fillna(0)
    df["breadth_30d"] = np.where(tot > 0, (df["up_30d"] - df["down_30d"]) / tot, np.nan)

    # Yahoo sometimes returns forwardEps and the eps_trend table on different bases.
    # If they disagree by >10%, G2, G3 and G4 are unreliable; only G1 (targets) is kept.
    conflict = (df["eps_fwd"].notna() & df["eps_ny_now"].notna()
                & ((df["eps_fwd"] / df["eps_ny_now"] - 1).abs() > 0.10))
    df["data_flag"] = np.where(conflict, "EPS sources disagree", "")
    df.loc[conflict, ["G2", "G3", "G4", "G4_30d", "eps_growth_ny"]] = np.nan

    for g in ("G1", "G2", "G3", "G4"):
        df.loc[~covered, g] = np.nan  # too few analysts: don't trust consensus

    num, den = np.zeros(len(df)), np.zeros(len(df))
    total_w = sum(r[4] for r in rules.values())
    for g, (_, d, bull, bear, w, neg) in rules.items():
        if mode == "Relative":
            s = score_relative(df[g], d, neg)
        else:
            s = df[g].apply(lambda v: score_value(v, d, bull, bear, neg))
        df[f"S_{g}"] = s
        has = s.notna().to_numpy()
        num += np.where(has, s.fillna(0) * w, 0)
        den += np.where(has, w, 0)
    df["score"] = np.where(den > 0, np.round(num / np.where(den > 0, den, 1) * 100, 1), np.nan)
    df["coverage"] = den / total_w if total_w else 0
    df["verdict"] = df["score"].apply(verdict)
    return df


def verdict(s) -> str:
    if pd.isna(s):
        return "No forward data"
    if s >= 40:
        return "Strong positive"
    if s >= 15:
        return "Positive"
    if s > -15:
        return "Neutral"
    if s > -40:
        return "Negative"
    return "Strong negative"


# Excel export: column -> (header, number format)
XL_COLS = {
    "symbol": ("Symbol", None), "name": ("Name", None), "section": ("Section", None),
    "sector": ("Sector", None), "price": ("Price ₹", "#,##0.0"), "mcap_cr": ("Mcap ₹cr", "#,##0"),
    "analysts": ("Analysts", "0"), "score": ("Score", "0"), "verdict": ("Verdict", None),
    "coverage": ("Coverage", "0%"), "G1": ("G1 Upside %", "0.0"), "G2": ("G2 FwdPE/TTM", "0.00"),
    "G3": ("G3 PEG", "0.00"), "G4": ("G4 Rev 90d %", "0.0"), "S_G1": ("G1 score", "0"),
    "S_G2": ("G2 score", "0"), "S_G3": ("G3 score", "0"), "S_G4": ("G4 score", "0"),
    "G4_30d": ("Rev 30d %", "0.0"), "breadth_30d": ("Up−down 30d", "0.00"), "pe_ttm": ("PE TTM", "0.0"),
    "pe_fwd": ("PE Fwd", "0.0"), "eps_growth_ny": ("Next-FY EPS gr %", "0.0"),
    "target_mean": ("Target mean", "#,##0"), "target_low": ("Target low", "#,##0"),
    "target_high": ("Target high", "#,##0"), "rec_key": ("Consensus", None), "data_flag": ("Data check", None),
}


def to_excel(view: pd.DataFrame, rules: dict, mode: str, min_analysts: int, data_date: str) -> bytes:
    """Formatted .xlsx: 'Scan' sheet (filtered table) + 'Settings' sheet (how it was scored)."""
    import io
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    hdr_fill = PatternFill("solid", fgColor="1F3A5F")
    body = Font(name="Arial", size=10)

    wb = Workbook()
    ws = wb.active
    ws.title = "Scan"
    cols = [c for c in XL_COLS if c in view.columns]
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j, value=XL_COLS[c][0])
        cell.font, cell.fill = hdr_font, hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, rec in enumerate(view[cols].itertuples(index=False), start=2):
        for j, (c, v) in enumerate(zip(cols, rec), start=1):
            if isinstance(v, float) and not np.isfinite(v):
                v = None
            elif isinstance(v, (np.floating, np.integer)):
                v = v.item()
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = body
            if XL_COLS[c][1]:
                cell.number_format = XL_COLS[c][1]
    last = max(len(view) + 1, 2)
    widths = {"name": 28, "section": 22, "sector": 20, "verdict": 15, "data_flag": 20, "rec_key": 12}
    for j, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(c, 11)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    if "score" in cols:
        L = get_column_letter(cols.index("score") + 1)
        ws.conditional_formatting.add(f"{L}2:{L}{last}", ColorScaleRule(
            start_type="num", start_value=-100, start_color="F8696B",
            mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=100, end_color="63BE7B"))
    for c in ("S_G1", "S_G2", "S_G3", "S_G4"):
        if c in cols:
            L = get_column_letter(cols.index(c) + 1)
            ws.conditional_formatting.add(f"{L}2:{L}{last}", CellIsRule(
                operator="equal", formula=["1"], fill=PatternFill("solid", fgColor="D9EAD3")))
            ws.conditional_formatting.add(f"{L}2:{L}{last}", CellIsRule(
                operator="equal", formula=["-1"], fill=PatternFill("solid", fgColor="F4CCCC")))

    st_ = wb.create_sheet("Settings")
    info = [
        ("Forward Sentiment Scanner — scan settings", None),
        ("Data date (snapshot)", data_date),
        ("Exported", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Stocks in this export", len(view)),
        ("Scoring mode", mode),
        ("Minimum analysts", min_analysts),
        ("", None),
    ]
    for i, (k, v) in enumerate(info, start=1):
        st_.cell(row=i, column=1, value=k).font = Font(name="Arial", size=12 if i == 1 else 10, bold=True)
        if v is not None:
            st_.cell(row=i, column=2, value=v).font = body
    r0 = len(info) + 1
    for j, h in enumerate(["Metric", "Description", "Better when", "Bull +1 (Absolute)",
                           "Bear −1 (Absolute)", "Weight", "≤ 0 scores −1"], start=1):
        cell = st_.cell(row=r0, column=j, value=h)
        cell.font, cell.fill = hdr_font, hdr_fill
    for i, (g, (label, d, bull, bear, w, neg)) in enumerate(rules.items(), start=r0 + 1):
        for j, v in enumerate([g, label, d, bull, bear, w, "Yes" if neg else "No"], start=1):
            st_.cell(row=i, column=j, value=v).font = body
    note_row = r0 + len(rules) + 2
    notes = [
        "Relative mode: each metric ranked within the scanned list — best third +1, middle 0, worst third −1.",
        "Absolute mode: fixed Bull/Bear thresholds above.",
        "Score = weighted average of metric scores with data, scaled to −100…+100.",
        "Verdict: ≥40 Strong positive · ≥15 Positive · −15 to 15 Neutral · ≤−15 Negative · ≤−40 Strong negative.",
        "Data: yfinance (unofficial Yahoo Finance). Rows flagged 'EPS sources disagree' keep only G1. "
        "Not investment advice.",
    ]
    for i, n in enumerate(notes, start=note_row):
        st_.cell(row=i, column=1, value=n).font = Font(name="Arial", size=9, italic=True)
    for col, wd in zip("ABCDEFG", [24, 34, 12, 16, 16, 9, 13]):
        st_.column_dimensions[col].width = wd

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def latest_snapshot() -> Path | None:
    snaps = sorted(SNAP_DIR.glob("forward_*.csv"))
    return snaps[-1] if snaps else None


# ----------------------------------------------------------------------------- UI
st.set_page_config(page_title="Forward Sentiment Scanner", page_icon="📈", layout="wide")
st.title("Forward Sentiment Scanner")
st.caption("Analyst targets, forward PE, PEG and EPS revisions from yfinance · NSE watchlist · not investment advice")

with st.sidebar:
    st.header("Stock list")
    up = st.file_uploader("Override list (CSV/TXT with a Symbol column)", type=["csv", "txt"])
    try:
        universe = load_universe(up) if up else load_universe(STOCKS_FILE)
    except FileNotFoundError:
        universe = pd.DataFrame(columns=["yf", "symbol", "section"])
        st.error("stocks.csv not found next to app.py.")
    symbols = universe["yf"].tolist()
    st.write(f"**{len(symbols)}** symbols loaded")

    st.header("Data")
    snap = latest_snapshot()
    today_file = SNAP_DIR / f"forward_{dt.date.today():%Y-%m-%d}.csv"
    st.write(f"Last snapshot: **{snap.stem.replace('forward_', '') if snap else 'none'}**")
    workers = st.slider("Parallel requests", 1, 8, 4, help="Lower this if Yahoo starts refusing requests (rate limit).")
    refresh = st.button("Fetch fresh data", type="primary", width="stretch",
                        help=f"About 3 requests per stock; {len(symbols)} stocks take a few minutes.")

    st.header("Scoring")
    mode = st.radio("Scoring mode", ["Relative", "Absolute"], horizontal=True,
                    help="Relative: each metric ranked within your list (top third +1, bottom third −1). "
                         "Absolute: fixed thresholds below. Analyst numbers are optimistic across the board, "
                         "so absolute mode rates most stocks positive.")
    min_analysts = st.number_input("Minimum analysts", 1, 20, 3)
    rules = {}
    with st.expander("Weights (and thresholds for Absolute mode)"):
        for g, (label, d, bull, bear, w, neg) in DEFAULT_RULES.items():
            st.markdown(f"**{g} · {label}** ({'higher' if d == 'Higher' else 'lower'} is better)")
            c1, c2, c3 = st.columns(3)
            b1 = c1.number_input("Bull +1", value=bull, key=f"{g}b", step=0.05)
            b2 = c2.number_input("Bear −1", value=bear, key=f"{g}r", step=0.05)
            ww = c3.number_input("Weight", value=w, key=f"{g}w", min_value=0, step=1)
            rules[g] = (label, d, b1, b2, ww, neg)

# ---- load or fetch
raw = None
if refresh and symbols:
    bar = st.progress(0.0, text="Starting…")
    raw = fetch_all(symbols, workers, bar)
    raw.to_csv(today_file, index=False)
    bar.empty()
    st.success(f"Fetched {len(raw)} stocks · saved snapshot {today_file.name}")
elif snap is not None:
    raw = pd.read_csv(snap)

if raw is None or raw.empty:
    st.info("No data yet. Click **Fetch fresh data** in the sidebar.")
    st.stop()

# keep only the current list and attach its sections
raw = raw.drop(columns=["section"], errors="ignore").merge(
    universe[["symbol", "section"]], on="symbol", how="inner")

df = compute(raw, rules, int(min_analysts), mode)

# ---- summary
c = st.columns(5)
c[0].metric("Stocks", len(df))
c[1].metric("With forward data", int(df["score"].notna().sum()))
c[2].metric("Positive or better", int((df["score"] >= 15).sum()))
c[3].metric("Negative or worse", int((df["score"] <= -15).sum()))
c[4].metric("Fetch errors", int((df.get("error", pd.Series(dtype=str)).fillna("") != "").sum()))

# ---- filters
sections = sorted({x for v in df["section"].fillna("") for x in v.split("; ") if x})
f0, f1, f2, f3 = st.columns([2, 2, 1, 2])
pick_sec = f0.multiselect("Section", sections, default=sections)
verdicts = ["Strong positive", "Positive", "Neutral", "Negative", "Strong negative", "No forward data"]
pick = f1.multiselect("Verdict", verdicts, default=verdicts[:5])
min_cov = f2.slider("Min data coverage", 0.0, 1.0, 0.5, 0.25)
q = f3.text_input("Search symbol or name")

in_sec = df["section"].fillna("").apply(lambda v: any(x in pick_sec for x in v.split("; ")) or not v)
view = df[in_sec & df["verdict"].isin(pick) & ((df["coverage"].fillna(0) >= min_cov) | df["score"].isna())]
if q:
    ql = q.lower()
    view = view[view["symbol"].str.lower().str.contains(ql) | view["name"].fillna("").str.lower().str.contains(ql)]
view = view.sort_values("score", ascending=False, na_position="last")

cols = ["symbol", "name", "section", "sector", "price", "mcap_cr", "analysts", "score", "verdict", "coverage",
        "G1", "G2", "G3", "G4", "G4_30d", "breadth_30d", "pe_ttm", "pe_fwd", "eps_growth_ny",
        "target_mean", "target_low", "target_high", "rec_key", "data_flag"]
st.dataframe(
    view[[c for c in cols if c in view]],
    hide_index=True, width="stretch", height=560,
    column_config={
        "symbol": st.column_config.TextColumn("Symbol"),
        "price": st.column_config.NumberColumn("Price ₹", format="%.1f"),
        "mcap_cr": st.column_config.NumberColumn("Mcap ₹cr", format="%.0f"),
        "analysts": st.column_config.NumberColumn("Analysts", format="%d"),
        "score": st.column_config.ProgressColumn("Score", min_value=-100, max_value=100, format="%.0f"),
        "coverage": st.column_config.NumberColumn("Coverage", format="percent"),
        "G1": st.column_config.NumberColumn("G1 Upside %", format="%.1f"),
        "G2": st.column_config.NumberColumn("G2 FwdPE/TTM", format="%.2f"),
        "G3": st.column_config.NumberColumn("G3 PEG", format="%.2f"),
        "G4": st.column_config.NumberColumn("G4 Rev 90d %", format="%.1f"),
        "G4_30d": st.column_config.NumberColumn("Rev 30d %", format="%.1f"),
        "breadth_30d": st.column_config.NumberColumn("Up−down 30d", format="%.2f",
                                                     help="(analysts up − down) ÷ total, next FY. yfinance counts are inconsistent; not scored."),
        "pe_ttm": st.column_config.NumberColumn("PE TTM", format="%.1f"),
        "pe_fwd": st.column_config.NumberColumn("PE Fwd", format="%.1f"),
        "eps_growth_ny": st.column_config.NumberColumn("Next-FY EPS gr %", format="%.1f"),
        "target_mean": st.column_config.NumberColumn("Target mean", format="%.0f"),
        "target_low": st.column_config.NumberColumn("Target low", format="%.0f"),
        "target_high": st.column_config.NumberColumn("Target high", format="%.0f"),
        "rec_key": st.column_config.TextColumn("Consensus"),
        "section": st.column_config.TextColumn("Section"),
        "data_flag": st.column_config.TextColumn("Data check"),
    },
)
data_date = today_file.stem.replace("forward_", "") if refresh else (snap.stem.replace("forward_", "") if snap else "")
dl1, dl2 = st.columns(2)
dl1.download_button(
    "Download Excel (.xlsx)",
    to_excel(view, rules, mode, int(min_analysts), data_date),
    f"forward_scan_{data_date or dt.date.today():}.xlsx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    type="primary", width="stretch",
)
dl2.download_button("Download CSV", view.to_csv(index=False).encode(), f"forward_scan_{data_date}.csv",
                    "text/csv", width="stretch")

# ---- single stock detail
st.subheader("Stock detail")
sel = st.selectbox("Choose a stock", view["symbol"].tolist() or df["symbol"].tolist())
if sel:
    r = df[df["symbol"] == sel].iloc[0]
    d1, d2 = st.columns(2)
    with d1:
        st.markdown(f"**{r.get('name') or sel}** · {r.get('sector') or ''}")
        st.write(f"Score **{r['score'] if pd.notna(r['score']) else '—'}** · {r['verdict']} · "
                 f"{int(r['analysts']) if pd.notna(r['analysts']) else 0} analysts")
        sc = pd.DataFrame({
            "Metric": [DEFAULT_RULES[g][0] for g in DEFAULT_RULES],
            "Value": [r[g] for g in DEFAULT_RULES],
            "Score": [r[f"S_{g}"] for g in DEFAULT_RULES],
        })
        st.dataframe(sc, hide_index=True, width="stretch")
    with d2:
        tr = pd.DataFrame({"Next-FY EPS consensus": ["90 days ago", "30 days ago", "Now"],
                           "₹": [r.get("eps_ny_90d"), r.get("eps_ny_30d"), r.get("eps_ny_now")]})
        st.dataframe(tr, hide_index=True, width="stretch")
        st.markdown(f"[Screener](https://www.screener.in/company/{sel}/consolidated/) · "
                    f"[TradingView](https://www.tradingview.com/chart/?symbol=NSE:{sel})")

errs = df[df.get("error", pd.Series("", index=df.index)).fillna("") != ""]
if not errs.empty:
    with st.expander(f"Fetch errors ({len(errs)})"):
        st.dataframe(errs[["symbol", "error"]], hide_index=True)

st.caption("Scores: each metric +1 / 0 / −1 (Relative: rank within your list; Absolute: fixed thresholds), weighted, rescaled to −100…+100 over metrics "
           "with data. G2 and G3 score −1 when zero or negative. Stocks with fewer analysts than the minimum get "
           "no forward score. Daily snapshots are kept in /snapshots so target and estimate changes can be "
           "tracked over time.")
